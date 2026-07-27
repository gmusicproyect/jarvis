"""Document loaders multi-formato."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import yaml

from jarvis.rag.base import RawDocument
from jarvis.utils.logging import get_logger

_LOG = get_logger("jarvis.rag.loaders")

CODE_EXTS = {
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".cs",
    ".java",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".swift",
    ".kt",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
    ".sql",
    ".sh",
    ".zsh",
}


class TextLoader:
    name = "text"
    exts = {".txt", ".md", ".markdown", ".rst", ".log"} | CODE_EXTS

    def can_load(self, path: Path) -> bool:
        return path.suffix.lower() in self.exts

    def load(self, path: Path) -> list[RawDocument]:
        text = path.read_text(encoding="utf-8", errors="replace")
        return [
            RawDocument(
                source=str(path.resolve()),
                text=text,
                metadata={"ext": path.suffix.lower(), "name": path.name},
            )
        ]


class JsonYamlLoader:
    name = "json_yaml"
    exts = {".json", ".yaml", ".yml"}

    def can_load(self, path: Path) -> bool:
        return path.suffix.lower() in self.exts

    def load(self, path: Path) -> list[RawDocument]:
        raw = path.read_text(encoding="utf-8", errors="replace")
        try:
            if path.suffix.lower() == ".json":
                data = json.loads(raw)
                text = json.dumps(data, ensure_ascii=False, indent=2)
            else:
                data = yaml.safe_load(raw)
                text = yaml.safe_dump(data, allow_unicode=True)
        except Exception:
            text = raw
        return [
            RawDocument(
                source=str(path.resolve()),
                text=text,
                metadata={"ext": path.suffix.lower(), "name": path.name},
            )
        ]


class CsvLoader:
    name = "csv"

    def can_load(self, path: Path) -> bool:
        return path.suffix.lower() == ".csv"

    def load(self, path: Path) -> list[RawDocument]:
        rows: list[str] = []
        with path.open(encoding="utf-8", errors="replace", newline="") as fh:
            reader = csv.reader(fh)
            for i, row in enumerate(reader):
                rows.append(f"fila {i+1}: " + " | ".join(row))
                if i > 5000:
                    rows.append("…(truncado)")
                    break
        return [
            RawDocument(
                source=str(path.resolve()),
                text="\n".join(rows),
                metadata={"ext": ".csv", "name": path.name},
            )
        ]


class PdfLoader:
    name = "pdf"

    def can_load(self, path: Path) -> bool:
        return path.suffix.lower() == ".pdf"

    def load(self, path: Path) -> list[RawDocument]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Instala pypdf para indexar PDF") from exc
        reader = PdfReader(str(path))
        docs: list[RawDocument] = []
        for i, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue
            docs.append(
                RawDocument(
                    source=str(path.resolve()),
                    text=text,
                    page=i,
                    metadata={"ext": ".pdf", "name": path.name, "page": str(i)},
                )
            )
        return docs


class DocxLoader:
    name = "docx"

    def can_load(self, path: Path) -> bool:
        return path.suffix.lower() == ".docx"

    def load(self, path: Path) -> list[RawDocument]:
        try:
            import docx  # python-docx
        except ImportError as exc:
            raise RuntimeError("Instala python-docx para indexar DOCX") from exc
        document = docx.Document(str(path))
        text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
        return [
            RawDocument(
                source=str(path.resolve()),
                text=text,
                metadata={"ext": ".docx", "name": path.name},
            )
        ]


class ExcelLoader:
    name = "excel"

    def can_load(self, path: Path) -> bool:
        return path.suffix.lower() in {".xlsx", ".xlsm"}

    def load(self, path: Path) -> list[RawDocument]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Instala openpyxl para indexar Excel") from exc
        wb = load_workbook(str(path), read_only=True, data_only=True)
        parts: list[str] = []
        for sheet in wb.worksheets:
            parts.append(f"# Hoja: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                vals = [str(c) if c is not None else "" for c in row]
                if any(vals):
                    parts.append(" | ".join(vals))
        return [
            RawDocument(
                source=str(path.resolve()),
                text="\n".join(parts),
                metadata={"ext": path.suffix.lower(), "name": path.name},
            )
        ]


class PptxLoader:
    name = "pptx"

    def can_load(self, path: Path) -> bool:
        return path.suffix.lower() == ".pptx"

    def load(self, path: Path) -> list[RawDocument]:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("Instala python-pptx para indexar PowerPoint") from exc
        prs = Presentation(str(path))
        docs: list[RawDocument] = []
        for i, slide in enumerate(prs.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    texts.append(shape.text)
            body = "\n".join(texts).strip()
            if body:
                docs.append(
                    RawDocument(
                        source=str(path.resolve()),
                        text=body,
                        page=i,
                        metadata={
                            "ext": ".pptx",
                            "name": path.name,
                            "page": str(i),
                        },
                    )
                )
        return docs


class CompositeLoader:
    """Selecciona el loader adecuado por extensión."""

    name = "composite"

    def __init__(self) -> None:
        self._loaders = [
            PdfLoader(),
            DocxLoader(),
            ExcelLoader(),
            PptxLoader(),
            CsvLoader(),
            JsonYamlLoader(),
            TextLoader(),
        ]

    def can_load(self, path: Path) -> bool:
        return any(loader.can_load(path) for loader in self._loaders)

    def load(self, path: Path) -> list[RawDocument]:
        for loader in self._loaders:
            if loader.can_load(path):
                try:
                    return loader.load(path)
                except Exception as exc:  # noqa: BLE001
                    _LOG.warning("load_failed", path=str(path), error=str(exc))
                    raise
        raise ValueError(f"Formato no soportado: {path.suffix}")
